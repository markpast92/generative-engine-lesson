import PyPDF2
import docx
import openpyxl
from io import BytesIO

def extract_from_pdf(file_bytes):
    """Estrae testo da PDF."""
    try:
        pdf_reader = PyPDF2.PdfReader(BytesIO(file_bytes))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
        return text.strip()
    except Exception as e:
        return f"Errore nell'estrazione PDF: {str(e)}"

def extract_from_docx(file_bytes):
    """Estrae testo da DOCX."""
    try:
        doc = docx.Document(BytesIO(file_bytes))
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text.strip()
    except Exception as e:
        return f"Errore nell'estrazione DOCX: {str(e)}"

def extract_from_xlsx(file_bytes):
    """Estrae testo da XLSX."""
    try:
        workbook = openpyxl.load_workbook(BytesIO(file_bytes))
        text = ""
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text += f"--- Sheet: {sheet_name} ---\n"
            for row in sheet.iter_rows(values_only=True):
                text += " | ".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
        return text.strip()
    except Exception as e:
        return f"Errore nell'estrazione XLSX: {str(e)}"

def extract_text(file_bytes, file_name):
    """Estrae testo dal file in base all'estensione."""
    file_extension = file_name.lower().split(".")[-1]
    
    if file_extension == "pdf":
        return extract_from_pdf(file_bytes)
    elif file_extension == "docx":
        return extract_from_docx(file_bytes)
    elif file_extension == "xlsx":
        return extract_from_xlsx(file_bytes)
    else:
        return f"Formato file non supportato: {file_extension}"